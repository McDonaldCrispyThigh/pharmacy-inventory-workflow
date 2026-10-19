"""Local structural/layout receipts and contact sheets for both document formats."""
from pathlib import Path
import json,re,hashlib
import fitz
from PIL import Image,ImageDraw
from docx import Document
from docx.oxml.ns import qn
import run_expanded_final as run
OUT,ROOT=run.OUT,run.ROOT
qa=OUT/'visual_qa';qa.mkdir(exist_ok=True)
records={}
for kind,path in [('pdf',ROOT/'docs/thesis/introduction_literature_methods_results_final_figures.pdf'),('word',OUT/'docx/plain_docx_review.pdf')]:
    dest=qa/kind;dest.mkdir(exist_ok=True)
    d=fitz.open(path);pages=[];overflow=[]
    for i,page in enumerate(d):
        text=page.get_text();png=dest/f'page_{i+1:02d}.png'
        page.get_pixmap(dpi=110).save(png)
        pages.append({'page':i+1,'text_characters':len(text),'tables':re.findall(r'Table\s+[34]\.\d',text),
            'figures':re.findall(r'Figure\s+[34]\.\d',text),'text':text})
        for block in page.get_text('dict')['blocks']:
            if block['type']!=0:continue
            for line in block['lines']:
                for span in line['spans']:
                    rect=fitz.Rect(span['bbox'])
                    boundary=page.rect*page.derotation_matrix
                    if rect.x0<boundary.x0-.5 or rect.y0<boundary.y0-.5 or rect.x1>boundary.x1+.5 or rect.y1>boundary.y1+.5:
                        overflow.append({'page':i+1,'text':span['text'],'bbox':list(rect)})
    for start in range(0,len(d),9):
        sheet=Image.new('RGB',(960,1290),'#dddddd');draw=ImageDraw.Draw(sheet)
        for j in range(start,min(start+9,len(d))):
            im=Image.open(dest/f'page_{j+1:02d}.png');im.thumbnail((306,396))
            col,row=(j-start)%3,(j-start)//3;left=col*320+7;top=row*430+25
            sheet.paste(im,(left,top));draw.text((left,top-19),f'{kind.upper()} page {j+1}',fill='black')
        sheet.save(dest/f'contact_{start+1:02d}.png')
    records[kind]={'pages':len(d),'file_sha256':run.sha(path),'page_boundary_violations':overflow,'page_index':pages}
    assert not overflow,overflow

path=OUT/'docx/From_Digital_Evidence_to_Pharmacy_Access_Introduction_through_Results.docx'
doc=Document(path);bad=[]
for elem in doc._element.xpath('.//w:r'):
    if not ''.join(elem.itertext()).strip():continue
    p=elem.rPr
    if p is None:continue
    for key in ('b','i'):
        e=p.find(qn('w:'+key))
        if e is not None and e.get(qn('w:val'),'1') not in ('0','false'):bad.append(key)
    c=p.find(qn('w:color'))
    if c is not None and c.get(qn('w:val')) not in ('000000','auto'):bad.append('color')
    z=p.find(qn('w:sz'))
    if z is not None and z.get(qn('w:val'))!='22':bad.append('size')
assert not bad,bad
assert len(doc.tables)==8 and len(doc._element.xpath('.//m:oMath'))==0
geometry=[]
for table in doc.tables:
    grid=[int(v.get(qn('w:w'))) for v in table._tbl.tblGrid.gridCol_lst]
    width=int(table._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w')))
    assert sum(grid)==width==9360
    for tr in table._tbl.tr_lst:
        assert sum(int(tc.tcPr.find(qn('w:tcW')).get(qn('w:w'))) for tc in tr.tc_lst)==9360
    geometry.append({'columns':len(grid),'width_twips':width,'rows':len(table.rows)})
formulas=run.read_json(OUT/'docx/formula_provenance.json')
assert sum(v['kind']=='DisplayMath' for v in formulas)==1
records['docx_structure']={'passed':True,'plain_black_11pt':True,'bold_italic_violations':0,'editable_tables':geometry,
    'word_equation_objects':0,'formal_equation_images':1,'unique_math_images':len(formulas),'figure_panels':5,
    'runtime_note':'The container runtime dependency loader is absent. Used the repository accepted local Pandoc/python-docx and hidden Word renderer. Skill render_docx.py could not start because pdf2image is absent; native Word rendered all pages successfully.'}
(qa/'document_structural_qa.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v.get('pages',v.get('editable_tables')) for k,v in records.items()}))

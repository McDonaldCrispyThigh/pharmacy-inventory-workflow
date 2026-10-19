"""Verify the stage-specific linkage diagnostic's thesis integration.

This script is read-only with respect to scientific inputs and writes only a QA
summary.  It does not execute the diagnostic, a provider workflow, or a model.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
from pathlib import Path

import pymupdf
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "docs/thesis"
CHAPTERS = THESIS / "chapters"
TABLES = THESIS / "tables"
FIGURES = THESIS / "figures/results"
DIAGNOSTIC = ROOT / "analysis/v2/board_linkage_stage_diagnostic_20260902T012636Z_v5"
FIGURE_PACKAGE = ROOT / "analysis/v2/final_results_20260902T012636Z_final_figure_integration_v1"
QA_PACKAGE = ROOT / "analysis/v2/stage_diagnostic_thesis_integration_20260903_v1"
PDF = THESIS / "introduction_literature_methods_results_final_figures.pdf"
LOG = THESIS / "introduction_literature_methods_results_final_figures.log"
LOT = THESIS / "introduction_literature_methods_results_final_figures.lot"
LOF = THESIS / "introduction_literature_methods_results_final_figures.lof"
TOC = THESIS / "introduction_literature_methods_results_final_figures.toc"
AUX = THESIS / "introduction_literature_methods_results_final_figures.aux"
WRAPPER = THESIS / "introduction_literature_methods_results_final_figures.tex"
DOCX = THESIS / "From_Digital_Evidence_to_Pharmacy_Access_Introduction_through_Results.docx"
WORD_PDF = (
    QA_PACKAGE
    / "docx_render_v4/From_Digital_Evidence_to_Pharmacy_Access_Introduction_through_Results.pdf"
)
OUTPUT = QA_PACKAGE / "qa/editorial_pass_qa.json"


UNCHANGED_HASHES = {
    CHAPTERS / "01_introduction_rewrite.tex": "57df4a99cc492401ed3c75d3eb4901df0f8625d3fc71a0abfe89f7d1bbcfd884",
    CHAPTERS / "02_literature_rewrite.tex": "aac053b46fdd0ce31bc0837e2352d2a7a8af80633287e056bae5fbbcef4b6b3b",
    TABLES / "table1_candidate_flow.tex": "51c48b5283cd3c9dff9db5a7e3bd26eed30ff160eb80ec5f878e4fa8e90ba831",
    TABLES / "table2_board_correspondence.tex": "670ab3b39ce2a64c7e616a14a72a252d53d3e90e8235899b2eab06749fbc74d7",
    TABLES / "table3_name_recurrence.tex": "87bd16ef342d91fa05218a28118199f19a1957369f0b29c361ce1fb06eb2a3f9",
    TABLES / "table4_paired_spatial.tex": "9d7f65b223b5a85cdf878f0c5b67314b0854a1a3840515084d5dfa20076457e7",
    TABLES / "table5_neighborhood.tex": "2ab0d91d13beb16b63013fef4453b6f3879a4e46c0df93874f8a0a67fe04f6b8",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def flat(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def pages_containing(document: pymupdf.Document, phrase: str) -> list[int]:
    return [index + 1 for index, page in enumerate(document) if phrase in flat(page.get_text())]


def footer_offsets(document: pymupdf.Document) -> list[float]:
    offsets: list[float] = []
    page_number = re.compile(r"^(?:[ivxlcdm]+|\d+)$", flags=re.IGNORECASE)
    for page_index, page in enumerate(document):
        if page_index == 0:
            continue
        candidates = []
        for block in page.get_text("rawdict")["blocks"]:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    chars = span.get("chars", [])
                    text = "".join(char["c"] for char in chars).strip()
                    if page_number.fullmatch(text) and span["bbox"][1] > page.rect.height - 70:
                        ink = [char["bbox"] for char in chars if not char["c"].isspace()]
                        require(bool(ink), f"Footer number has no ink on page {page_index + 1}")
                        candidates.append(
                            (
                                min(box[0] for box in ink),
                                min(box[1] for box in ink),
                                max(box[2] for box in ink),
                                max(box[3] for box in ink),
                            )
                        )
        require(len(candidates) == 1, f"Expected one footer number on Word PDF page {page_index + 1}")
        bbox = candidates[0]
        offsets.append(((bbox[0] + bbox[2]) / 2) - page.rect.width / 2)
    return offsets


def main() -> None:
    required = [
        PDF,
        LOG,
        LOT,
        LOF,
        TOC,
        AUX,
        WRAPPER,
        DOCX,
        WORD_PDF,
        DIAGNOSTIC / "artifact_manifest.json",
        DIAGNOSTIC / "board_linkage_stage_diagnostic.json",
        FIGURE_PACKAGE / "final_figure_manifest.json",
        FIGURES / "Figure_3.1_workflow_asset_manifest.json",
    ]
    for path in required:
        require(path.is_file(), f"Missing required artifact: {path}")

    artifact_manifest = json.loads(
        (DIAGNOSTIC / "artifact_manifest.json").read_text(encoding="utf-8")
    )
    require(artifact_manifest["artifact_count"] == 22, "Unexpected v5 artifact count")
    for name, expected in artifact_manifest["artifacts"].items():
        require(sha256(DIAGNOSTIC / name) == expected, f"v5 artifact hash mismatch: {name}")

    diagnostic = json.loads(
        (DIAGNOSTIC / "board_linkage_stage_diagnostic.json").read_text(encoding="utf-8")
    )
    require(all(diagnostic["verification_checks"].values()), "A v5 verification check failed")
    require(not diagnostic["provenance"]["provider_calls"], "v5 reports a provider call")
    require(not diagnostic["provenance"]["model_calls"], "v5 reports a model call")
    require(diagnostic["provenance"]["classification_changes"] == 0, "v5 changed classifications")

    frame = diagnostic["comparison_frame"]
    place = diagnostic["place_id_analysis"]
    locations = diagnostic["outlet_location_unit_analysis"]
    require(
        [
            frame["raw_candidates_scored_before_geographic_reporting_filter"],
            frame["raw_candidates_in_primary_60_zip_frame"],
            frame["raw_candidates_outside_primary_60_zip_frame"],
            frame["included_candidates_in_primary_frame"],
            frame["excluded_candidates_in_primary_frame"],
            frame["screened_board_records_in_primary_frame"],
        ]
        == [1226, 984, 242, 339, 645, 178],
        "Core v5 frame counts differ",
    )
    require(
        [row["count"] for row in place["candidate_cross_tab"]] == [212, 18, 127, 627],
        "Candidate cross-tab differs",
    )
    require(
        [row["count"] for row in place["board_four_stage_decomposition"]] == [27, 5, 0, 146],
        "Board-stage decomposition differs",
    )
    require(
        [locations["raw_place_ids"], locations["raw_collapsed_locations"], locations["included_place_ids"], locations["included_collapsed_locations"]]
        == [984, 772, 339, 239],
        "Location-unit counts differ",
    )

    arithmetic = {
        "984-339": 984 - 339 == 645,
        "151-146": 151 - 146 == 5,
        "645/984": math.isclose(round(645 / 984 * 100, 2), 65.55),
        "146/151": math.isclose(round(146 / 151 * 100, 2), 96.69),
        "5/151": math.isclose(round(5 / 151 * 100, 2), 3.31),
        "5/178": math.isclose(round(5 / 178 * 100, 2), 2.81),
        "212/230": math.isclose(round(212 / 230 * 100, 2), 92.17),
        "127/754": math.isclose(round(127 / 754 * 100, 2), 16.84),
        "board_stage_sum": 27 + 5 + 0 + 146 == 178,
        "candidate_cross_tab_sum": 212 + 18 + 127 + 627 == 984,
    }
    require(all(arithmetic.values()), "New arithmetic failed")

    for path, expected in UNCHANGED_HASHES.items():
        require(sha256(path) == expected, f"Previously approved source changed: {path.name}")

    methods = (CHAPTERS / "03_methodology.tex").read_text(encoding="utf-8")
    results = (CHAPTERS / "04_results.tex").read_text(encoding="utf-8")
    table = (TABLES / "table3_stage_specific_filtering.tex").read_text(encoding="utf-8")
    wrapper = WRAPPER.read_text(encoding="utf-8")
    methods_flat = flat(methods)
    results_flat = flat(results)
    table_flat = flat(table)
    require(
        methods_flat.index("Edges below 75")
        < methods_flat.index("A post hoc stage-specific diagnostic")
        < methods_flat.index("The primary analysis retains each unique provider Place ID"),
        "Methods diagnostic placement differs",
    )
    require(
        "Classifications and eligibility decisions were frozen before this downstream analysis"
        in methods_flat,
        "Methods does not document frozen decisions",
    )
    require(
        "neither Board information nor diagnostic results were returned to the model or reducer"
        in methods_flat,
        "Methods does not document the no-feedback boundary",
    )
    require(
        "effect of the LLM/reducer eligibility stage" in methods_flat
        and "retained by the LLM/reducer eligibility stage" in methods_flat,
        "Methods does not use the standardized stage term",
    )
    rq2_opening = (
        "RQ2 first establishes the screened Board denominator and primary 60-ZIP linkage, "
        "then examines the effect of the LLM/reducer eligibility stage, unmatched workflow "
        "records, name-recurrence patterns, and the predefined geographic-frame sensitivity."
    )
    require(rq2_opening in results_flat, "RQ2 opening sentence differs")
    require(
        results_flat.index("Within the primary 60-ZIP frame")
        < results_flat.index("In a post hoc location-unit sensitivity"),
        "Results diagnostic placement differs",
    )
    required_results = [
        "covered 151 of the 178 screened Board records",
        "removed 645 candidate rows",
        "retaining 146 of the 151 Board records linked before filtering under the same rule",
        "212 were included and 18 were excluded",
        "127 were included and 627 were excluded",
        "92.17\\% of candidates with at least one qualified Board edge were included",
        "16.84\\% of candidates without a qualified edge",
        "These percentages are descriptive enrichment measures under the declared linkage rule, not classification-performance estimates",
        "No Board record had a qualified included edge but was lost solely through the one-to-one assignment",
        "The 96.69\\% Board-referenced retention was unchanged after exact normalized address--ZIP collapse",
        "the raw and included location frames covered 151 and 146 Board records, respectively",
        "rather than general LLM classification accuracy",
    ]
    require(all(item in results_flat for item in required_results), "A Results statement is missing")
    require(
        results.index(r"\input{tables/table2_board_correspondence}")
        < results.index(r"\input{tables/table3_stage_specific_filtering}")
        < results.index(r"\input{tables/table3_name_recurrence}"),
        "Results table order differs",
    )
    table_tokens = [
        "LLM/reducer eligibility and Board correspondence in the primary 60-ZIP frame",
        "Include & 212 & 127 & 339",
        "Exclude & 18 & 627 & 645",
        "All in-frame candidates & 230 & 754 & 984",
        "No qualified raw-candidate edge & 27 & 15.17",
        "Qualified raw edge, but all endpoints excluded & 5 & 2.81",
        "Qualified included edge, but no final assignment & 0 & 0.00",
        "Final assigned link & 146 & 82.02",
        "The five Board records associated with attrition at the LLM/reducer eligibility stage equal 3.31",
        "151 Board records covered by the raw-candidate assignment",
    ]
    require(all(token in table_flat for token in table_tokens), "New table content differs")

    thesis_sources = list(CHAPTERS.glob("*.tex"))
    require(
        not any("\\subsection{" in path.read_text(encoding="utf-8") for path in thesis_sources),
        "An x.x.x heading is present",
    )
    require("05_discussion" not in wrapper and "06_conclusion" not in wrapper, "Later chapter is active")
    require("\\abstract{" not in wrapper, "A results-bearing abstract is active")

    lot = LOT.read_text(encoding="utf-8")
    lof = LOF.read_text(encoding="utf-8")
    toc = TOC.read_text(encoding="utf-8")
    aux = AUX.read_text(encoding="utf-8")
    expected_tables = [
        "Candidate flow and forced-choice decisions",
        "Primary Board correspondence, linkage specifications, and location-unit sensitivity",
        "LLM/reducer eligibility and Board correspondence in the primary 60-ZIP frame",
        "Name recurrence and primary Board correspondence",
        "Paired tract classification by inventory and distance threshold",
        "Neighborhood-quartile potential-proximity patterns",
    ]
    require(lot.count("\\contentsline {table}") == 6, "List of Tables does not have six entries")
    require(all(title in lot for title in expected_tables), "List of Tables title missing")
    require(
        all(f"\\numberline {{{number}}}" in lot for number in ["4.1", "4.2", "4.3", "4.4", "4.5", "4.6"]),
        "List of Tables numbering differs",
    )
    require(lof.count("\\contentsline {figure}") == 4, "List of Figures does not have four entries")
    require("\\contentsline {subsection}" not in toc, "Contents contains an x.x.x heading")
    expected_label_numbers = {
        "tab:candidate_flow": "4.1",
        "tab:board_correspondence": "4.2",
        "tab:stage_specific_filtering": "4.3",
        "tab:name_recurrence": "4.4",
        "tab:paired_spatial": "4.5",
        "tab:neighborhood_patterns": "4.6",
    }
    require(
        all(f"\\newlabel{{{label}}}{{{{{number}}}" in aux for label, number in expected_label_numbers.items()),
        "A LaTeX table label has an unexpected number",
    )
    require("??" not in aux, "Unresolved LaTeX reference marker")
    log = LOG.read_text(encoding="utf-8", errors="replace")
    forbidden_log = [
        "Overfull \\hbox",
        "Overfull \\vbox",
        "Underfull \\hbox",
        "Underfull \\vbox",
        "undefined references",
        "multiply defined",
        "LaTeX Warning:",
        "! Undefined control sequence",
        "! LaTeX Error",
    ]
    require(not any(token.lower() in log.lower() for token in forbidden_log), "LaTeX log warning")

    result_manifest = json.loads(
        (FIGURE_PACKAGE / "final_figure_manifest.json").read_text(encoding="utf-8")
    )
    figure_hashes: dict[str, str] = {}
    for record in result_manifest["figures"]:
        stem = record["name"]
        for suffix, metadata in record["formats"].items():
            path = FIGURES / f"{stem}.{suffix}"
            actual = sha256(path)
            require(actual == metadata["sha256"], f"Approved Results figure changed: {path.name}")
            figure_hashes[path.name] = actual
    workflow_manifest = json.loads(
        (FIGURES / "Figure_3.1_workflow_asset_manifest.json").read_text(encoding="utf-8")
    )
    for name, expected in workflow_manifest["sha256"].items():
        path = FIGURES / name
        actual = sha256(path)
        require(actual == expected, f"Workflow figure changed: {name}")
        figure_hashes[name] = actual

    with pymupdf.open(PDF) as document:
        latex_pages = {
            "methods": pages_containing(document, "A post hoc stage-specific diagnostic isolated"),
            "results_start": pages_containing(document, "Within the primary 60-ZIP frame"),
            "results_candidate": pages_containing(document, "At the candidate level, 230 raw candidates"),
            "results_location": pages_containing(
                document, "The 96.69% Board-referenced retention was unchanged"
            ),
            "table": pages_containing(
                document, "LLM/reducer eligibility and Board correspondence"
            )[-1:],
        }
        require(all(len(pages) == 1 for pages in latex_pages.values()), "LaTeX placement is ambiguous")
        table_page = latex_pages["table"][0]
        require(
            "Final assigned link" in flat(document[table_page - 1].get_text()),
            "LaTeX Table 4.3 is split or incomplete",
        )
        require("??" not in "".join(page.get_text() for page in document), "LaTeX PDF has unresolved marker")

    word = Document(DOCX)
    require(len(word.tables) == 6, "Word does not contain six tables")
    require(len(word.inline_shapes) == 5, "Word does not contain the five approved image objects")
    require(
        not any(paragraph.style.name == "Heading 3" for paragraph in word.paragraphs),
        "Word contains an x.x.x heading",
    )
    word_text = flat(" ".join(paragraph.text for paragraph in word.paragraphs))
    require("Table 4.3." in word_text, "Word is missing Table 4.3")
    require("Table 4.4; Figure 4.1" in word_text, "Word name-recurrence reference is stale")
    require(
        all(f"Table 4.{number}. {title}" in word_text for number, title in enumerate(expected_tables, 1)),
        "A Word table caption or number differs",
    )
    require("[tab:" not in word_text and "??" not in word_text, "Word contains unresolved reference text")

    with pymupdf.open(WORD_PDF) as document:
        word_pages = {
            "methods": pages_containing(document, "A post hoc stage-specific diagnostic isolated"),
            "results_start": pages_containing(document, "Within the primary 60-ZIP frame"),
            "results_candidate": pages_containing(document, "At the candidate level, 230 raw candidates"),
            "results_location": pages_containing(
                document, "The 96.69% Board-referenced retention was unchanged"
            ),
            "table": pages_containing(
                document, "LLM/reducer eligibility and Board correspondence"
            )[-1:],
        }
        require(all(len(pages) == 1 for pages in word_pages.values()), "Word placement is ambiguous")
        table_page = word_pages["table"][0]
        require(
            "Final assigned link" in flat(document[table_page - 1].get_text()),
            "Word Table 4.3 is split or incomplete",
        )
        offsets = footer_offsets(document)
        require(max(abs(value) for value in offsets) <= 1.0, "Word page numbers are not centered")

    summary = {
        "status": "pass",
        "scope": "thesis integration of frozen v5 stage-specific diagnostic",
        "provider_calls": False,
        "model_calls": False,
        "diagnostic_rerun": False,
        "frozen_classification_changes": 0,
        "linkage_rule_changes": 0,
        "spatial_analysis_changes": 0,
        "v5_artifact_hashes_verified": 22,
        "arithmetic": arithmetic,
        "unchanged_approved_source_hashes": {
            str(path.relative_to(ROOT)).replace("\\", "/"): expected
            for path, expected in UNCHANGED_HASHES.items()
        },
        "figure_hashes": figure_hashes,
        "heading_gate": {"latex_x_x_x": 0, "word_heading_3": 0},
        "lists": {"tables": 6, "logical_figures": 4},
        "latex_physical_pages": latex_pages,
        "word_physical_pages": word_pages,
        "word_footer_max_center_offset_points": max(abs(value) for value in offsets),
        "latex_pdf_sha256": sha256(PDF),
        "word_docx_sha256": sha256(DOCX),
        "word_qa_pdf_sha256": sha256(WORD_PDF),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()

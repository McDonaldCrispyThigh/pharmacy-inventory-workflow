"""Use the accepted plain DOCX workflow in a new isolated output folder."""
import argparse,importlib.util,sys,fitz
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'analysis/maintenance/methods_reproducibility_20260907_v1/build_docx.py'
spec=importlib.util.spec_from_file_location('accepted_plain_builder',SOURCE)
base=importlib.util.module_from_spec(spec);spec.loader.exec_module(base)
base.WORK=ROOT/'analysis/expanded_final/20260908_v1/docx'
base.INPUT=base.WORK/'build_input';base.OUTPUT=base.WORK/base.legacy.FINAL_DOCX.name

def prepare():
    base.prepare()
    # The accepted preparation already renders panel A. Refresh panel B from its
    # current vector source as well, so an older cached PNG cannot enter Word.
    name='Figure_3.1B_correspondence_proximity_matched_scale'
    d=fitz.open(base.THESIS/'figures/results'/f'{name}.pdf')
    d[0].get_pixmap(dpi=300).save(base.INPUT/'figures/results'/f'{name}.png');d.close()

def build():
    base.build()
    from docx import Document
    from docx.oxml.ns import qn
    d=Document(base.OUTPUT)
    names=base.legacy.FIGURE_TITLES
    alt=[
      'Base 60-ZIP, three-term retrieval, 60-result saturation check, bounded subdivision and Place-ID union precede frozen Details and direct Wayback evidence. Development-only pilot did not validate corrected archive packets. Board-blind validation and reduction form the inventory.',
      'Fixed 60-ZIP linkage and fixed 99-ZIP Board supply; all included workflow coordinates enter the same 327-origin proximity analysis. RQ2 links do not filter RQ3 supply.',
      'Repeated and single-occurrence workflow correspondence: 45.49 and 19.86 percent. Board correspondence: 93.13 and 59.57 percent. Intervals use 5,000 component resamples.',
      'Aligned maps of 327 tract classifications. At 0.5 and 1 mile, 48 and 46 tracts disagree between inventories. Blue indicates workflow only within; amber Board only within.',
      'Signed workflow-minus-Board distance histogram. Median minus 0.018 miles; first and third quartiles minus 0.099 and 0.015 miles. Fixed 0.1-mile bins.'
    ]
    import json
    bindings=json.loads((base.WORK/'image_bindings.json').read_text())
    i=0
    for shape,b in zip(d.inline_shapes,bindings):
        if b['figure']:
            shape._inline.docPr.set('title',names[i]);shape._inline.docPr.set('descr',alt[i]);i+=1
    assert i==5
    # Keep each editable table with its explanatory note; the prior builder could
    # leave that note alone at the top of the following page.
    for table in d.tables:
        for cell in table.rows[-1].cells:
            for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
    for paragraph in d.paragraphs:
        if paragraph.text.startswith('Note.'):
            paragraph.paragraph_format.keep_together=True
    d.save(base.OUTPUT)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('stage',choices=['prepare','build','render']);v=p.parse_args().stage
    {'prepare':prepare,'build':build,'render':base.render}[v]()

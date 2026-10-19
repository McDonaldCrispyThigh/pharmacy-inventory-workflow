"""Build the stage-gated Introduction-through-Results thesis DOCX.

The script converts the current LaTeX chapter sources with Pandoc, then applies a
Word-native thesis layout with python-docx.  It does not run any scientific analysis
and does not alter the approved figure files.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import uuid
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "docs" / "thesis"
WORK = ROOT / "analysis" / "docx_work"
BUILD_INPUT = WORK / "build_input"
RAW_DOCX = WORK / "thesis_pandoc.docx"
PREWORD_DOCX = WORK / "thesis_preword_current.docx"
FINAL_DOCX = (
    THESIS
    / "From_Digital_Evidence_to_Pharmacy_Access_Introduction_through_Results.docx"
)
PANDOC = Path.home() / "AppData" / "Local" / "Pandoc" / "pandoc.exe"

TITLE = (
    "From Digital Evidence to Pharmacy Access:\n"
    "An Evidence-Guided Establishment Data Workflow"
)

FIGURE_LIST_ENTRIES = [
    ("Figure 3.1.", "Evidence-guided pharmacy establishment workflow."),
    ("Figure 4.1.", "Primary correspondence by normalized-name recurrence."),
    ("Figure 4.2.", "Paired spatial disagreement among 327 tract origins."),
    ("Figure 4.3.", "Workflow-minus-Board nearest-distance differences across 327 tract origins."),
]

FIGURE_ALT_TEXTS = [
    "Flowchart of the Board-blind inventory-construction stage. A fixed 60-ZIP, "
    "three-term query plan produces deduplicated Google Place IDs; Details metadata "
    "and returned reviews join selected Wayback evidence in a fixed order. A pre-run "
    "A/B/C pilot precedes forced-choice classification, deterministic validation, "
    "and reduction to the workflow pharmacy inventory; Board and downstream "
    "variables are withheld.",
    "Flowchart of downstream correspondence and proximity analyses. The workflow "
    "inventory and screened Board inventory enter deterministic 60-ZIP record "
    "linkage, producing matched and unmatched records. Their eligible outlet "
    "coordinates separately enter a common projected nearest-distance analysis for "
    "327 tract origins; linkage matches do not filter the paired 0.5-, 1-, and "
    "2-mile proximity comparison.",
    "Dot-and-interval plot comparing correspondence for repeated-name and "
    "single-occurrence components in the workflow and Board inventories.",
    "Two tract maps showing the four paired spatial states at 0.5- and 1-mile "
    "thresholds in Hennepin and Ramsey Counties.",
    "Histogram of workflow-minus-Board nearest-distance differences across "
    "327 tract origins.",
]

FIGURE_TITLES = [
    "Figure 3.1, panel (a)",
    "Figure 3.1, panel (b)",
    "Figure 4.1",
    "Figure 4.2",
    "Figure 4.3",
]

TABLE_CAPTIONS = [
    "Candidate flow and forced-choice decisions",
    "Primary Board correspondence, linkage specifications, and location-unit sensitivity",
    "LLM/reducer eligibility and Board correspondence in the primary 60-ZIP frame",
    "Name recurrence and primary Board correspondence",
    "Paired tract classification by inventory and distance threshold",
    "Neighborhood-quartile potential-proximity patterns",
]


def safe_reset(directory: Path) -> None:
    resolved = directory.resolve()
    work_resolved = WORK.resolve()
    if resolved.parent != work_resolved or resolved.name != "build_input":
        raise RuntimeError(f"Refusing to reset unexpected directory: {resolved}")
    if resolved.exists():
        shutil.rmtree(resolved)
    resolved.mkdir(parents=True)


def prepare_pandoc_input() -> Path:
    WORK.mkdir(parents=True, exist_ok=True)
    safe_reset(BUILD_INPUT)
    (BUILD_INPUT / "chapters").mkdir()
    (BUILD_INPUT / "tables").mkdir()
    (BUILD_INPUT / "figures" / "results").mkdir(parents=True)

    wrapper = THESIS / "introduction_literature_methods_results_final_figures.tex"
    shutil.copy2(wrapper, BUILD_INPUT / wrapper.name)
    shutil.copy2(THESIS / "references.bib", BUILD_INPUT / "references.bib")

    chapter_names = [
        "01_introduction_rewrite.tex",
        "02_literature_rewrite.tex",
        "03_methodology.tex",
        "04_results.tex",
    ]
    for name in chapter_names:
        src = THESIS / "chapters" / name
        text = src.read_text(encoding="utf-8")
        if name == "03_methodology.tex":
            # Pandoc's TeX-math reader accepts these equivalent labels reliably.
            text = text.replace(r"\mathrm{name}", r"\mathit{name}")
            text = text.replace(r"\mathrm{address}", r"\mathit{address}")
            text = text.replace(r"\mathrm{geography}", r"\mathit{geography}")
            text = text.replace(
                "figures/results/Figure_3.1A_inventory_construction_thesis.pdf",
                "figures/results/Figure_3.1A_inventory_construction_thesis.png",
            )
            text = text.replace(
                "figures/results/Figure_3.1B_correspondence_proximity_matched_scale.pdf",
                "figures/results/Figure_3.1B_correspondence_proximity_matched_scale.png",
            )
            # Keep the two workflow panels as ordinary sequential figures in the
            # DOCX-only conversion. Word styling below turns the second caption
            # into an unlisted continuation, so Figure 3.1 appears only once in
            # the List of Figures.
            text = text.replace(r"\textbf{(a)}\par\medskip", "")
            text = text.replace(r"\textbf{(b)}\par\medskip", "")
            text = text.replace(r"\makeatletter", "")
            text = text.replace(r"\makeatother", "")
            text = text.replace(
                r"\@makecaption{Figure~\ref{fig:workflow_overview} (continued)}{",
                r"\caption{Figure 3.1 (continued). ",
            )
        if name == "04_results.tex":
            text = text.replace("figures/results/figure01_name_recurrence.pdf", "figures/results/figure01_name_recurrence.png")
            text = text.replace("figures/results/figure02_spatial_disagreement.pdf", "figures/results/figure02_spatial_disagreement.png")
            text = text.replace("figures/results/figure03_signed_distance_histogram.pdf", "figures/results/figure03_signed_distance_histogram.png")
            # Caption paragraphs are later rewritten as plain Word text so their
            # list-of-figures entries can stay concise. Preserve negative values
            # through that rewrite by using a typographic minus rather than an
            # inline TeX math object in the DOCX-only build copy.
            text = text.replace("$-1.909$", "−1.909")
            text = text.replace("$-0.009$", "−0.009")
            text = text.replace("$-0.078$", "−0.078")
            text = text.replace("$-0.072$", "−0.072")
        (BUILD_INPUT / "chapters" / name).write_text(text, encoding="utf-8")

    for src in sorted((THESIS / "tables").glob("table*.tex")):
        text = src.read_text(encoding="utf-8")
        text = text.replace(r"\resizebox{\textwidth}{!}{%", "")
        text = re.sub(r"\\end\{tabular\}%\s*\n\}\s*", r"\\end{tabular}\n", text)
        text = re.sub(r"\\begin\{minipage\}\{[^\n]+\}", "", text)
        text = text.replace(r"\end{minipage}", "")
        text = text.replace(r"\footnotesize", "")
        text = re.sub(r"(?m)^\\small\s*$", "", text)
        text = text.replace(r"\begin{table}[p]", r"\begin{table}[htbp]")
        (BUILD_INPUT / "tables" / src.name).write_text(text, encoding="utf-8")

    for src in sorted((THESIS / "figures" / "results").glob("*.png")):
        shutil.copy2(src, BUILD_INPUT / "figures" / "results" / src.name)

    return BUILD_INPUT / wrapper.name


def run_pandoc(source: Path) -> None:
    if not PANDOC.exists():
        raise FileNotFoundError(f"Pandoc not found: {PANDOC}")
    cmd = [
        str(PANDOC),
        source.name,
        "--from=latex",
        "--to=docx",
        "--standalone",
        "--citeproc",
        "--bibliography=references.bib",
        "--metadata=reference-section-title:References",
        "--resource-path=.;chapters;tables;figures;figures/results",
        f"--output={RAW_DOCX}",
    ]
    subprocess.run(cmd, cwd=BUILD_INPUT, check=True)


def set_font(run, name: str, size: float | None = None, bold=None, italic=None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name: str, size: float, bold=None, italic=None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def get_or_add_style(doc: Document, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)


def set_keep_lines(paragraph, keep_next=False) -> None:
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = True
    if keep_next:
        paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, instruction: str, placeholder: str = "Update field in Word") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hidden_field(paragraph, instruction: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    field.append(OxmlElement("w:r"))
    paragraph._p.append(field)


def insert_front_matter(first_heading) -> None:
    p = first_heading.insert_paragraph_before("Contents")
    p.style = "Front Matter Title"
    toc = first_heading.insert_paragraph_before()
    toc.style = "TOC Holder"
    add_field(toc, r'TOC \o "1-2" \h \z \u')

    p = first_heading.insert_paragraph_before("List of Tables")
    p.style = "Front Matter Title"
    p.paragraph_format.page_break_before = True
    lot = first_heading.insert_paragraph_before()
    lot.style = "TOC Holder"
    add_field(lot, r'TOC \h \z \t "Table Caption,1"')

    p = first_heading.insert_paragraph_before("List of Figures")
    p.style = "Front Matter Title"
    p.paragraph_format.page_break_before = True
    lof = first_heading.insert_paragraph_before()
    lof.style = "TOC Holder"
    add_field(lof, r'TOC \h \z \t "Figure Caption,1"')


def replace_xml_text(doc: Document, replacements: dict[str, str]) -> None:
    for node in doc.element.xpath(".//w:t"):
        if not node.text:
            continue
        for old, new in replacements.items():
            node.text = node.text.replace(old, new)


def set_cell_margins(cell, top=55, start=70, bottom=55, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        # CT_TcPr requires cell margins before vertical alignment. Appending the
        # element after w:vAlign produces a package that Word may try to repair
        # indefinitely, so insert it at the schema-defined position.
        v_align = tc_pr.find(qn("w:vAlign"))
        if v_align is None:
            tc_pr.append(tc_mar)
        else:
            tc_pr.insert(tc_pr.index(v_align), tc_mar)
    # Use the broadly supported left/right names in the stored Word package.
    # The start/end parameters remain part of the Python API for call-site clarity.
    for margin, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        # CT_TblPr places borders before layout, look, and caption metadata.
        following = next(
            (
                tbl_pr.find(qn(f"w:{tag}"))
                for tag in ("tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription")
                if tbl_pr.find(qn(f"w:{tag}")) is not None
            ),
            None,
        )
        if following is None:
            tbl_pr.append(borders)
        else:
            tbl_pr.insert(tbl_pr.index(following), borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")
    for edge in ("top", "bottom"):
        node = borders.find(qn(f"w:{edge}"))
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "000000")

    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.first_child_found_in("w:tcBorders")
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = tc_borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            tc_borders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")


def set_row_bottom_border(row, value: str, size: str = "6") -> None:
    """Set one row's lower rule without changing other table borders."""

    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.first_child_found_in("w:tcBorders")
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = tc_borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            tc_borders.append(bottom)
        bottom.set(qn("w:val"), value)
        if value != "nil":
            bottom.set(qn("w:sz"), size)
            bottom.set(qn("w:color"), "000000")


def set_table_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    # Presence alone means true and avoids version-specific ST_OnOff values.
    header.attrib.pop(qn("w:val"), None)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def apply_table_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = Inches(widths[min(idx, len(widths) - 1)])
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is not None:
                tc_w.set(qn("w:w"), str(int(width.twips)))
                tc_w.set(qn("w:type"), "dxa")


def format_tables(doc: Document) -> None:
    widths = [
        [1.30, 3.05, 0.90, 1.00],
        [2.95, 0.65, 1.00, 0.90, 0.90],
        [2.80, 1.35, 1.35, 0.70],
        [0.80, 2.15, 0.65, 0.65, 0.75, 1.35],
        [0.75, 0.80, 0.80, 1.10, 1.00, 0.80, 0.80],
        [0.75, 0.55, 0.40, 1.15, 1.15, 0.80, 0.75, 0.75],
    ]
    sizes = [9.0, 8.2, 8.3, 8.5, 8.5, 8.0]
    for index, table in enumerate(doc.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_fixed_layout(table)
        set_table_borders(table)
        apply_table_widths(table, widths[index])
        repeat_header(table.rows[0])
        if index == 2:
            # The stage-specific diagnostic has two internal panels. Keep the
            # panel label clear of the column-header rule and separate each
            # panel's header from its data rows.
            set_row_bottom_border(table.rows[0], "nil")
            set_row_bottom_border(table.rows[2], "single")
            set_row_bottom_border(table.rows[7], "single")
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                for p in cell.paragraphs:
                    if index == 2:
                        p.alignment = (
                            WD_ALIGN_PARAGRAPH.LEFT
                            if cell_index == 0
                            else WD_ALIGN_PARAGRAPH.CENTER
                        )
                    else:
                        p.alignment = (
                            WD_ALIGN_PARAGRAPH.LEFT
                            if index in (0, 1, 3, 5)
                            else WD_ALIGN_PARAGRAPH.CENTER
                        )
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.paragraph_format.left_indent = Inches(0)
                    p.paragraph_format.right_indent = Inches(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    # A repeating header should not remain alone at a page foot;
                    # keep it with the first data row when the table begins.
                    set_keep_lines(p, keep_next=(row_index == 0))
                    for run in p.runs:
                        is_header = row_index == 0 or (
                            index == 2 and row_index in (1, 2, 6, 7)
                        )
                        set_font(run, "Times New Roman", sizes[index], bold=is_header)


def format_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, "Times New Roman", 11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Inches(0.30)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.widow_control = True

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, "Times New Roman", 15, bold=True)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(24)
    h1.paragraph_format.line_spacing = 1.0
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, "Times New Roman", 12, bold=True)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.line_spacing = 1.0
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, "Times New Roman", 11, bold=True)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.line_spacing = 1.0
    h3.paragraph_format.keep_with_next = True

    chapter = get_or_add_style(doc, "Chapter Label")
    set_style_font(chapter, "Times New Roman", 15, bold=True)
    chapter.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chapter.paragraph_format.space_before = Pt(18)
    chapter.paragraph_format.space_after = Pt(18)
    chapter.paragraph_format.line_spacing = 1.0
    chapter.paragraph_format.keep_with_next = True

    front = get_or_add_style(doc, "Front Matter Title")
    set_style_font(front, "Times New Roman", 24, bold=True)
    front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    front.paragraph_format.space_before = Pt(18)
    front.paragraph_format.space_after = Pt(18)
    front.paragraph_format.line_spacing = 1.0
    front.paragraph_format.keep_with_next = True

    toc_holder = get_or_add_style(doc, "TOC Holder")
    set_style_font(toc_holder, "Times New Roman", 11)
    toc_holder.paragraph_format.first_line_indent = Inches(0)
    toc_holder.paragraph_format.line_spacing = 1.0
    toc_holder.paragraph_format.space_after = Pt(0)

    for level, indent in ((1, 0.0), (2, 0.25), (3, 0.50)):
        toc_style = get_or_add_style(doc, f"TOC {level}")
        set_style_font(toc_style, "Times New Roman", 11, bold=(level == 1))
        toc_style.paragraph_format.left_indent = Inches(indent)
        toc_style.paragraph_format.first_line_indent = Inches(0)
        toc_style.paragraph_format.space_before = Pt(0)
        toc_style.paragraph_format.space_after = Pt(2)
        toc_style.paragraph_format.line_spacing = 1.0

    table_caption = get_or_add_style(doc, "Table Caption")
    set_style_font(table_caption, "Times New Roman", 9.5)
    table_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_caption.paragraph_format.first_line_indent = Inches(0)
    table_caption.paragraph_format.space_before = Pt(10)
    table_caption.paragraph_format.space_after = Pt(5)
    table_caption.paragraph_format.line_spacing = 1.0
    table_caption.paragraph_format.keep_with_next = True

    figure_caption = get_or_add_style(doc, "Figure Caption")
    set_style_font(figure_caption, "Times New Roman", 9.5)
    figure_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_caption.paragraph_format.first_line_indent = Inches(0)
    figure_caption.paragraph_format.space_before = Pt(5)
    figure_caption.paragraph_format.space_after = Pt(8)
    figure_caption.paragraph_format.line_spacing = 1.0
    figure_caption.paragraph_format.keep_together = True

    figure_continuation = get_or_add_style(doc, "Figure Continuation Caption")
    set_style_font(figure_continuation, "Times New Roman", 9.5)
    figure_continuation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_continuation.paragraph_format.first_line_indent = Inches(0)
    figure_continuation.paragraph_format.space_before = Pt(5)
    figure_continuation.paragraph_format.space_after = Pt(8)
    figure_continuation.paragraph_format.line_spacing = 1.0
    figure_continuation.paragraph_format.keep_together = True

    figure_panel_marker = get_or_add_style(doc, "Figure Panel Marker")
    set_style_font(figure_panel_marker, "Times New Roman", 10, bold=True)
    figure_panel_marker.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_panel_marker.paragraph_format.first_line_indent = Inches(0)
    figure_panel_marker.paragraph_format.space_before = Pt(0)
    figure_panel_marker.paragraph_format.space_after = Pt(4)
    figure_panel_marker.paragraph_format.line_spacing = 1.0
    figure_panel_marker.paragraph_format.keep_with_next = True

    table_note = get_or_add_style(doc, "Table Note")
    set_style_font(table_note, "Times New Roman", 9)
    table_note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_note.paragraph_format.first_line_indent = Inches(0)
    table_note.paragraph_format.space_before = Pt(4)
    table_note.paragraph_format.space_after = Pt(8)
    table_note.paragraph_format.line_spacing = 1.0

    bibliography = get_or_add_style(doc, "Bibliography")
    set_style_font(bibliography, "Times New Roman", 10)
    bibliography.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bibliography.paragraph_format.first_line_indent = Inches(-0.30)
    bibliography.paragraph_format.left_indent = Inches(0.30)
    bibliography.paragraph_format.space_before = Pt(0)
    bibliography.paragraph_format.space_after = Pt(5)
    bibliography.paragraph_format.line_spacing = 1.0

    equation = get_or_add_style(doc, "Equation")
    set_style_font(equation, "Cambria Math", 11)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.first_line_indent = Inches(0)
    equation.paragraph_format.space_before = Pt(6)
    equation.paragraph_format.space_after = Pt(6)
    equation.paragraph_format.line_spacing = 1.0
    equation.paragraph_format.keep_together = True


def format_title_page(doc: Document, first_heading_index: int) -> None:
    title_paras = doc.paragraphs[:first_heading_index]
    for p in title_paras:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            set_font(run, "Times New Roman", 11)

    title_paras[0].paragraph_format.space_before = Pt(112)
    title_paras[0].paragraph_format.space_after = Pt(18)
    for run in title_paras[0].runs:
        set_font(run, "Times New Roman", 14, bold=True)

    title_paras[1].paragraph_format.space_after = Pt(6)
    title_paras[2].paragraph_format.space_after = Pt(104)
    for run in title_paras[2].runs:
        set_font(run, "Times New Roman", 12, bold=True)

    title_paras[3].paragraph_format.space_after = Pt(36)
    title_paras[4].paragraph_format.space_after = Pt(34)
    title_paras[5].paragraph_format.space_after = Pt(0)


def number_and_format_body(doc: Document) -> None:
    chapter_no = 0
    section_no = 0
    subsection_no = 0
    references_started = False
    original_paragraphs = list(doc.paragraphs)

    for p in original_paragraphs:
        style_name = p.style.name
        text = p.text.strip()
        if style_name == "Heading 1":
            if text == "References":
                references_started = True
                p.paragraph_format.page_break_before = True
                continue
            chapter_no += 1
            section_no = 0
            subsection_no = 0
            label = p.insert_paragraph_before(f"Chapter {chapter_no}")
            label.style = "Chapter Label"
            if chapter_no > 1:
                label.paragraph_format.page_break_before = True
            continue

        if references_started:
            if text:
                p.style = "Bibliography"
            continue

        if style_name == "Heading 2":
            section_no += 1
            subsection_no = 0
            p.text = f"{chapter_no}.{section_no}  {text}"
        elif style_name == "Heading 3":
            subsection_no += 1
            p.text = f"{chapter_no}.{section_no}.{subsection_no}  {text}"
        elif style_name == "Table Caption":
            p.style = "Table Caption"
        elif style_name == "Image Caption":
            p.style = "Figure Caption"
        elif style_name in {"Front Matter Title", "TOC Holder", "Chapter Label"}:
            continue
        elif text.startswith("Note."):
            p.style = "Table Note"
        elif p._p.xpath(".//m:oMath | .//m:oMathPara") and not text:
            p.style = "Equation"
        elif text:
            p.style = "Normal"

        set_keep_lines(p, keep_next=style_name in ("Heading 1", "Heading 2", "Heading 3"))


def prefix_captions(doc: Document) -> None:
    table_index = 0
    figure_index = 0
    for p in doc.paragraphs:
        if p.style.name == "Table Caption":
            table_index += 1
            if not p.text.startswith("Table "):
                p.text = f"Table 4.{table_index}. {p.text}"
        elif p.style.name == "Figure Caption":
            figure_index += 1
            if figure_index == 1:
                if not p.text.startswith("Figure "):
                    p.text = f"Figure 3.1. {p.text}"
            elif figure_index == 2:
                if not p.text.startswith("Figure 3.1 (continued)."):
                    p.text = f"Figure 3.1 (continued). {p.text}"
                p.style = "Figure Continuation Caption"
            else:
                results_index = figure_index - 2
                if not p.text.startswith("Figure "):
                    p.text = f"Figure 4.{results_index}. {p.text}"
    if table_index != 6 or figure_index != 5:
        raise RuntimeError(
            f"Expected six table captions and five figure-panel captions; "
            f"got {table_index} and {figure_index}"
        )


def format_figures(doc: Document) -> None:
    widths = [
        Inches(6.50),
        Inches(6.50),
        Inches(6.20),
        Inches(8.55),
        Inches(6.20),
    ]
    if len(doc.inline_shapes) != 5:
        raise RuntimeError(f"Expected five inline figures, found {len(doc.inline_shapes)}")
    for index, (shape, width) in enumerate(zip(doc.inline_shapes, widths), 1):
        ratio = shape.height / shape.width
        shape.width = width
        shape.height = int(width * ratio)
        p = shape._inline.getparent().getparent().getparent()
        p_obj = next(par for par in doc.paragraphs if par._p is p)
        p_obj.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_obj.paragraph_format.first_line_indent = Inches(0)
        p_obj.paragraph_format.space_before = Pt(6)
        p_obj.paragraph_format.space_after = Pt(0)
        p_obj.paragraph_format.line_spacing = 1.0
        p_obj.paragraph_format.keep_with_next = True
        p_obj.paragraph_format.keep_together = True
        if index in (1, 2):
            marker = p_obj.insert_paragraph_before("(a)" if index == 1 else "(b)")
            marker.style = "Figure Panel Marker"
            marker.paragraph_format.page_break_before = True
            marker.paragraph_format.keep_with_next = True

    section_32 = next(
        (
            p
            for p in doc.paragraphs
            if p.style.name == "Heading 2"
            and p.text.strip().endswith("Candidate Retrieval and Evidence Collection")
        ),
        None,
    )
    if section_32 is None:
        raise RuntimeError("Could not find Methods section 3.2 after workflow figures")
    section_32.paragraph_format.page_break_before = True


def protect_display_equations(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        if paragraph.style.name != "Equation":
            continue
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = True
        if index > 0:
            paragraphs[index - 1].paragraph_format.keep_with_next = True


def protect_lead_ins(doc: Document) -> None:
    """Keep short introductory lines with the material they introduce."""

    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "Three research questions guide the study:":
            paragraph.paragraph_format.keep_with_next = True


def compact_pilot_condition_list(doc: Document) -> None:
    """Keep the A/B/C evidence conditions as three compact Word paragraphs."""

    labels = ("A.", "B.", "C.")
    for label in labels:
        paragraphs = list(doc.paragraphs)
        matches = [p for p in paragraphs if p.text.strip() == label]
        if len(matches) != 1:
            raise RuntimeError(
                f"Expected one standalone pilot-condition label {label!r}; "
                f"found {len(matches)}"
            )
        label_paragraph = matches[0]
        index = paragraphs.index(label_paragraph)
        if index + 1 >= len(paragraphs):
            raise RuntimeError(f"Pilot-condition label {label!r} has no text")
        text_paragraph = paragraphs[index + 1]
        item_text = text_paragraph.text.strip()
        if not item_text:
            raise RuntimeError(f"Pilot-condition label {label!r} has empty text")
        label_paragraph.text = f"{label} {item_text}"
        label_paragraph.style = "Normal"
        label_paragraph.paragraph_format.left_indent = Inches(0.30)
        label_paragraph.paragraph_format.first_line_indent = Inches(0)
        label_paragraph.paragraph_format.keep_together = True
        text_paragraph._element.getparent().remove(text_paragraph._element)


def build_preword() -> Path:
    source = prepare_pandoc_input()
    run_pandoc(source)
    doc = Document(RAW_DOCX)
    if len(doc.tables) != 6:
        raise RuntimeError(f"Pandoc conversion produced {len(doc.tables)} tables instead of six")
    if len(doc.inline_shapes) != 5:
        raise RuntimeError(f"Pandoc conversion produced {len(doc.inline_shapes)} figures instead of five")

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    first_heading_index = next(
        i for i, p in enumerate(doc.paragraphs) if p.style.name == "Heading 1"
    )
    format_styles(doc)
    format_title_page(doc, first_heading_index)
    first_heading = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
    insert_front_matter(first_heading)
    number_and_format_body(doc)
    prefix_captions(doc)
    format_tables(doc)
    format_figures(doc)
    protect_lead_ins(doc)
    compact_pilot_condition_list(doc)
    protect_display_equations(doc)

    replace_xml_text(
        doc,
        {
            "[tab:board_correspondence]": "4.2",
            "[tab:stage_specific_filtering]": "4.3",
            "[tab:name_recurrence]": "4.4",
            "[tab:paired_spatial]": "4.5",
            "[tab:neighborhood_patterns]": "4.6",
            "[fig:workflow_overview]": "3.1",
            "[fig:name_recurrence_final]": "4.1",
            "[fig:spatial_disagreement_final]": "4.2",
            "[fig:signed_distance_final]": "4.3",
        },
    )

    # Reapply the style-level fonts to runs created during caption and heading edits.
    for p in doc.paragraphs:
        if p.style.name in {
            "Heading 1",
            "Heading 2",
            "Heading 3",
            "Chapter Label",
            "Figure Panel Marker",
        }:
            for run in p.runs:
                set_font(run, "Times New Roman", p.style.font.size.pt, bold=True)
        elif p.style.name in {
            "Table Caption",
            "Figure Caption",
            "Figure Continuation Caption",
        }:
            for run in p.runs:
                set_font(run, "Times New Roman", 9.5)

    core_props = doc.core_properties
    core_props.title = TITLE.replace("\n", " ")
    core_props.author = "Congyuan (Othello) Zheng"
    core_props.subject = "Introduction through Results thesis draft"
    core_props.keywords = "pharmacy establishments; record linkage; potential proximity"
    core_props.comments = "Generated from the current stage-gated LaTeX thesis sources."

    PREWORD_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(PREWORD_DOCX)
    return PREWORD_DOCX


def word_find_paragraph(document, prefix: str):
    for paragraph in document.Paragraphs:
        text = paragraph.Range.Text.replace("\r", "").replace("\x07", "").strip()
        if text.startswith(prefix):
            return paragraph
    raise RuntimeError(f"Word paragraph not found: {prefix}")


def word_find_styled_paragraph(document, prefix: str, style_name: str):
    """Find a body paragraph without accidentally matching a TOC result row."""

    for paragraph in document.Paragraphs:
        text = paragraph.Range.Text.replace("\r", "").replace("\x07", "").strip()
        if text.startswith(prefix) and paragraph.Style.NameLocal == style_name:
            return paragraph
    raise RuntimeError(f"Word paragraph not found: {prefix} [{style_name}]")


def clear_story_range(story_range) -> None:
    story_range.Text = ""


def add_centered_page_number(section, number_style: int, restart: bool) -> None:
    footer = section.Footers.Item(1)
    footer.LinkToPrevious = False
    clear_story_range(footer.Range)
    footer.Range.ParagraphFormat.Alignment = 1
    footer.PageNumbers.RestartNumberingAtSection = restart
    if restart:
        footer.PageNumbers.StartingNumber = 1
    footer.PageNumbers.NumberStyle = number_style
    footer.Range.Fields.Add(footer.Range, -1, "PAGE", True)
    # Inserting the field can inherit Word's default left alignment even when
    # the empty footer range was centered immediately beforehand. Reapply the
    # alignment and clear inherited body indents after field creation so the
    # saved document and fixed-format export agree geometrically.
    footer_format = footer.Range.ParagraphFormat
    footer_format.LeftIndent = 0
    footer_format.RightIndent = 0
    footer_format.FirstLineIndent = 0
    footer_format.Alignment = 1


def finalize_with_word(preword: Path, output: Path, qa_pdf: Path) -> Path:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    word.AutomationSecurity = 3
    word.Options.SaveNormalPrompt = False
    document = None
    temp_output = output.with_name(f".{output.stem}.{uuid.uuid4().hex}.docx")
    temp_pdf = qa_pdf.with_name(f".{qa_pdf.stem}.{uuid.uuid4().hex}.pdf")
    try:
        print("word:open", flush=True)
        document = word.Documents.Open(
            str(preword.resolve()),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
            Revert=True,
            NoEncodingDialog=True,
            OpenAndRepair=True,
        )
        print("word:sections", flush=True)

        contents = word_find_paragraph(document, "Contents")
        rng = document.Range(contents.Range.Start, contents.Range.Start)
        rng.InsertBreak(2)  # wdSectionBreakNextPage

        chapter_one = word_find_paragraph(document, "Chapter 1")
        rng = document.Range(chapter_one.Range.Start, chapter_one.Range.Start)
        rng.InsertBreak(2)

        # The pre-Word build places a page-break-before panel marker ahead of
        # each workflow image. Begin the detailed Methods sequence on a fresh
        # page after the second panel.
        methods_32 = word_find_styled_paragraph(
            document,
            "3.2  Candidate Retrieval and Evidence Collection",
            "Heading 2",
        )
        methods_32.Format.PageBreakBefore = True

        # Keep the two compact Results tables from beginning late on a page and
        # carrying only their final panel or a split header onto the next one.
        # These are layout-only controls; the table content is unchanged.
        table_one = word_find_styled_paragraph(document, "Table 4.1.", "Table Caption")
        table_one.Format.PageBreakBefore = True
        table_one.Format.KeepWithNext = True
        table_three = word_find_styled_paragraph(document, "Table 4.3.", "Table Caption")
        table_three.Format.PageBreakBefore = True
        table_three.Format.KeepWithNext = True

        # Isolate the approved metropolitan map on a landscape page.
        map_shape = document.InlineShapes.Item(4)
        map_start = map_shape.Range.Paragraphs.Item(1).Range.Start
        document.Range(map_start, map_start).InsertBreak(2)
        map_caption = word_find_paragraph(document, "Figure 4.2.")
        document.Range(map_caption.Range.End, map_caption.Range.End).InsertBreak(2)
        map_caption = word_find_paragraph(document, "Figure 4.2.")
        map_section = map_caption.Range.Sections.Item(1)
        map_section.PageSetup.Orientation = 1  # wdOrientLandscape
        map_section.PageSetup.TopMargin = 72
        map_section.PageSetup.BottomMargin = 72
        map_section.PageSetup.LeftMargin = 72
        map_section.PageSetup.RightMargin = 72
        map_index = map_section.Index
        if map_index < document.Sections.Count:
            following = document.Sections.Item(map_index + 1)
            following.PageSetup.Orientation = 0  # wdOrientPortrait
            following.PageSetup.TopMargin = 72
            following.PageSetup.BottomMargin = 72
            following.PageSetup.LeftMargin = 72
            following.PageSetup.RightMargin = 72

        # Title page unnumbered; front matter lower-case Roman; body Arabic.
        for section in document.Sections:
            section.Headers.Item(1).LinkToPrevious = False
            clear_story_range(section.Headers.Item(1).Range)
            section.PageSetup.HeaderDistance = 35.4
            section.PageSetup.FooterDistance = 35.4

        for footer_kind in (1, 2, 3):
            try:
                document.Sections.Item(1).Footers.Item(footer_kind).LinkToPrevious = False
                clear_story_range(document.Sections.Item(1).Footers.Item(footer_kind).Range)
            except Exception:
                pass

        add_centered_page_number(document.Sections.Item(2), 2, True)
        add_centered_page_number(document.Sections.Item(3), 0, True)
        for idx in range(4, document.Sections.Count + 1):
            footer = document.Sections.Item(idx).Footers.Item(1)
            footer.LinkToPrevious = True
            footer.PageNumbers.RestartNumberingAtSection = False

        # Linking a later section can reset the apparent paragraph alignment in
        # Word's footer story. Apply centering to every numbered section after
        # all links have been established.
        for idx in range(2, document.Sections.Count + 1):
            footer_format = (
                document.Sections.Item(idx).Footers.Item(1).Range.ParagraphFormat
            )
            footer_format.LeftIndent = 0
            footer_format.RightIndent = 0
            footer_format.FirstLineIndent = 0
            footer_format.Alignment = 1

        # Update Contents, List of Tables, and List of Figures. The helper builds
        # each list from the final body text, then shortens the four logical
        # List-of-Figures hyperlinks while retaining their live PAGEREF fields.
        print("word:update-toc", flush=True)
        update_short_figure_list(document)
        document.Repaginate()

        output.parent.mkdir(parents=True, exist_ok=True)
        qa_pdf.parent.mkdir(parents=True, exist_ok=True)
        print("word:save", flush=True)
        document.SaveAs2(str(temp_output.resolve()), FileFormat=16)
        print("word:export", flush=True)
        document.ExportAsFixedFormat(str(temp_pdf.resolve()), 17)
        print("word:close", flush=True)
        document.Close(SaveChanges=False)
        document = None
        temp_output.replace(output)
        temp_pdf.replace(qa_pdf)
    finally:
        if document is not None:
            document.Close(SaveChanges=False)
        word.Quit()
        pythoncom.CoUninitialize()
    return output


def replace_once_in_word(document, old: str, new: str) -> None:
    """Replace one expected phrase while failing closed on stale source text."""

    occurrences = document.Content.Text.count(old)
    if occurrences == 0 and document.Content.Text.count(new) == 1:
        return
    if occurrences != 1:
        raise RuntimeError(
            f"Expected exactly one Word occurrence of {old!r}; found {occurrences}"
        )
    search_range = document.Content.Duplicate
    finder = search_range.Find
    finder.ClearFormatting()
    finder.Text = old
    finder.Forward = True
    finder.Wrap = 0
    finder.Format = False
    if not finder.Execute():
        raise RuntimeError(f"Word did not replace expected phrase: {old!r}")
    search_range.Text = new
    if document.Content.Text.count(new) != 1:
        raise RuntimeError(f"Word did not retain replacement phrase: {new!r}")


def update_short_figure_list(document) -> None:
    """Refresh TOCs, then shorten only the displayed List-of-Figures entries."""

    document.Repaginate()
    for toc in document.TablesOfContents:
        toc.Update()
    if document.TablesOfContents.Count != 3:
        raise RuntimeError(
            f"Expected Contents, List of Tables, and List of Figures; "
            f"found {document.TablesOfContents.Count} TOCs"
        )
    figure_list = document.TablesOfContents.Item(3)
    expected_labels = [label for label, _ in FIGURE_LIST_ENTRIES]
    entries = []
    for index in range(1, figure_list.Range.Paragraphs.Count + 1):
        paragraph = figure_list.Range.Paragraphs.Item(index)
        text = paragraph.Range.Text.replace("\r", "").strip()
        if any(text.startswith(label) for label in expected_labels):
            entries.append(paragraph)
    if len(entries) != len(FIGURE_LIST_ENTRIES):
        raise RuntimeError(
            f"Expected {len(FIGURE_LIST_ENTRIES)} List-of-Figures entries; "
            f"found {len(entries)} linked entries"
        )
    for index, ((label, short_caption), entry) in enumerate(
        zip(FIGURE_LIST_ENTRIES, entries), 1
    ):
        entry_text = entry.Range.Text.replace("\r", "").strip()
        if not entry_text.startswith(label):
            raise RuntimeError(
                f"List-of-Figures entry {index} does not start with {label!r}: "
                f"{entry_text!r}"
            )
        if entry.Range.Hyperlinks.Count < 1:
            raise RuntimeError(
                f"Expected a hyperlink in List-of-Figures entry {index}; "
                f"found {entry.Range.Hyperlinks.Count}"
            )
        hyperlink = entry.Range.Hyperlinks.Item(1)
        anchor = hyperlink.SubAddress
        if not anchor:
            raise RuntimeError(f"List-of-Figures entry {index} has no bookmark target")
        hyperlink.TextToDisplay = f"{label} {short_caption}"

        # Re-add the tab leader and live page-reference field removed when Word
        # replaces the hyperlink's display text.
        tail = document.Range(entry.Range.End - 1, entry.Range.End - 1)
        tail.Text = "\t"
        tail.Collapse(0)  # wdCollapseEnd
        document.Fields.Add(tail, -1, f"PAGEREF {anchor} \\h", True)
        entry.Range.Font.Color = 0  # wdColorBlack
        entry.Range.Font.Underline = 0  # wdUnderlineNone


def repair_existing_with_word(source: Path, output: Path, qa_pdf: Path) -> Path:
    """Apply final Word-native QA fixes to the already validated DOCX."""

    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    word.AutomationSecurity = 3
    word.Options.SaveNormalPrompt = False
    document = None
    temp_output = output.with_name(f".{output.stem}.{uuid.uuid4().hex}.docx")
    temp_pdf = qa_pdf.with_name(f".{qa_pdf.stem}.{uuid.uuid4().hex}.pdf")
    try:
        print("word:open-existing", flush=True)
        document = word.Documents.Open(
            str(source.resolve()),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
            Revert=True,
            NoEncodingDialog=True,
            OpenAndRepair=True,
        )

        replace_once_in_word(
            document,
            "Before retrieval or linkage, the query plan fixes 60 ZIP codes in . Postal",
            "Before retrieval or linkage, the query plan fixes 60 ZIP codes in "
            "code/study_design.py. Postal",
        )
        replace_once_in_word(
            document,
            "before any Google linkage",
            "before record linkage",
        )

        # Keep the first Results table together on a fresh page.
        table_one = word_find_styled_paragraph(document, "Table 4.1.", "Table Caption")
        table_one.Format.PageBreakBefore = True
        table_one.Format.KeepWithNext = True

        # Keep the compact two-panel stage diagnostic from beginning with a
        # partial header at the foot of the preceding page.
        table_three = word_find_styled_paragraph(document, "Table 4.3.", "Table Caption")
        table_three.Format.PageBreakBefore = True
        table_three.Format.KeepWithNext = True

        # Keep each displayed equation with its introducing paragraph.
        equation_count = 0
        for index in range(1, document.Paragraphs.Count + 1):
            paragraph = document.Paragraphs.Item(index)
            if paragraph.Style.NameLocal != "Equation":
                continue
            equation_count += 1
            paragraph.Format.KeepTogether = True
            if index > 1:
                document.Paragraphs.Item(index - 1).Format.KeepWithNext = True
        if equation_count != 2:
            raise RuntimeError(f"Expected two displayed equations; found {equation_count}")

        # Set descriptive alternatives without changing the embedded image bytes.
        if document.InlineShapes.Count != len(FIGURE_ALT_TEXTS):
            raise RuntimeError(
                f"Expected {len(FIGURE_ALT_TEXTS)} inline figures; "
                f"found {document.InlineShapes.Count}"
            )
        for index, alt_text in enumerate(FIGURE_ALT_TEXTS, 1):
            shape = document.InlineShapes.Item(index)
            shape.AlternativeText = alt_text
            shape.Title = FIGURE_TITLES[index - 1]

        print("word:update-lists", flush=True)
        update_short_figure_list(document)

        output.parent.mkdir(parents=True, exist_ok=True)
        qa_pdf.parent.mkdir(parents=True, exist_ok=True)
        print("word:save-existing", flush=True)
        document.SaveAs2(str(temp_output.resolve()), FileFormat=16)
        print("word:export-existing", flush=True)
        document.ExportAsFixedFormat(str(temp_pdf.resolve()), 17)
        document.Close(SaveChanges=False)
        document = None
        temp_output.replace(output)
        temp_pdf.replace(qa_pdf)
    finally:
        if document is not None:
            document.Close(SaveChanges=False)
        word.Quit()
        pythoncom.CoUninitialize()
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--finalize-word", action="store_true")
    parser.add_argument("--repair-existing", action="store_true")
    parser.add_argument("--output", type=Path, default=FINAL_DOCX)
    parser.add_argument(
        "--qa-pdf", type=Path, default=WORK / "word_export" / "thesis_word_render.pdf"
    )
    args = parser.parse_args()

    if args.finalize_word and args.repair_existing:
        parser.error("Choose either --finalize-word or --repair-existing")
    if args.repair_existing:
        if not args.output.exists():
            raise FileNotFoundError(f"Existing Word document not found: {args.output}")
        path = repair_existing_with_word(args.output, args.output, args.qa_pdf)
    elif args.finalize_word:
        if not PREWORD_DOCX.exists():
            raise FileNotFoundError(f"Build the pre-Word document first: {PREWORD_DOCX}")
        path = finalize_with_word(PREWORD_DOCX, args.output, args.qa_pdf)
    else:
        path = build_preword()
    print(path)


if __name__ == "__main__":
    main()
